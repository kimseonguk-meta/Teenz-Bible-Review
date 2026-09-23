"""Build the Matthew 3-way comparison docx from FIXED patches: MSG vs Teenz EN(수정) vs Teenz KO(수정)."""
import json, re, os, glob, subprocess
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from msg_parse import parse_msg_sections, range_to_set

WORK = '/home/hatch/workspace/teenz-bible-review'
FIX = os.path.join(WORK, 'fixes')
OUT_DIR = os.path.expanduser('~/workspace/your_files/teenz-bible-review-matthew')
os.makedirs(OUT_DIR, exist_ok=True)


def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)


def set_cell_text(cell, text, bold=False, size=9, color=None, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def add_para(cell, text, bold=False, size=9, color=None, italic=False, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p


def norm(t):
    return re.sub(r'\s+', ' ', t.strip().lower())


def load_patches(prefix):
    out = {}
    for path in glob.glob(os.path.join(FIX, f'{prefix}_*.json')):
        for c in json.load(open(path, encoding='utf-8')):
            out[c['chapter']] = c
    return out


def load_msg():
    msg = {}
    raw_path = os.path.join(WORK, 'msg_matthew_raw.json')
    if os.path.exists(raw_path):
        msg.update(json.load(open(raw_path, encoding='utf-8')))
    for ch in ('1', '3'):
        p = os.path.join(WORK, f'msg_raw_{ch}.txt')
        if ch not in msg and os.path.exists(p):
            msg[ch] = open(p, encoding='utf-8').read()
    return msg


def main():
    en_p = load_patches('en_matthew')
    ko_p = load_patches('ko_matthew')
    assert len(en_p) == 28 and len(ko_p) == 28, f'patches incomplete: EN {len(en_p)}, KO {len(ko_p)}'
    msg = load_msg()
    mt_verse_counts = [25, 23, 17, 25, 48, 34, 29, 34, 38, 42, 30, 50, 58, 36, 39, 28, 27, 35, 30, 34, 46, 46, 39, 51, 46, 75, 66, 20]

    # 현재 앱 데이터 (수정 전 비교용)
    app_en = {c['num']: c for c in json.load(open('/home/hatch/workspace/teenz-fix/client/src/data/allBibleData.json', encoding='utf-8'))['Matthew']}
    ko_raw = open('/home/hatch/workspace/teenz-fix/client/src/data/gospelDataKo.ts', encoding='utf-8').read()
    mstart = ko_raw.find('"Matthew": [')
    mend = ko_raw.find('"Mark": [', mstart)
    ko_block = ko_raw[mstart:mend]
    app_ko = {}
    for part in re.split(r'\{\s*\n\s*"num": ', ko_block)[1:]:
        num = int(part.split(',')[0])
        pm = re.search(r'"paragraphs": \[(.*?)\]', part, re.S)
        paras = re.findall(r'"((?:[^"\\]|\\.)*)"', pm.group(1)) if pm else []
        app_ko[num] = paras
    old_en_sets = {n: set(norm(p) for p in c['paragraphs']) for n, c in app_en.items()}
    old_ko_sets = {n: set(norm(p) for p in ps) for n, ps in app_ko.items()}

    doc = Document()
    for section in doc.sections:
        section.orientation = 1
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(1.2)
        section.top_margin = section.bottom_margin = Cm(1.2)
    doc.styles['Normal'].font.size = Pt(10)

    doc.add_heading('Teenz Bible 번역 대조 — 마태복음 (수정본)', level=1)
    intro = doc.add_paragraph(
        '2026-09-22 확정된 번역 4대 원칙(MSG 원문 기준 · MSG 절 구분 · 의도적 요약 금지 · EN/KO 구조 일치)을 적용하여 '
        'Teenz EN/KO를 수정한 결과물입니다. 각 행은 The Message의 절 구간 기준이며, '
        '이번에 새로 복원·수정된 문단은 ✚ 표시와 연두색 배경으로 구분했습니다. '
        '부록의 ‘수정 내역’과 ‘컨펌 필요 목록’을 먼저 보시면 전체 변경점을 빠르게 파악할 수 있습니다.'
    )
    intro.runs[0].font.size = Pt(10)
    doc.add_paragraph('작성: 2026-09-22 · 앱에는 미반영 (성욱 컨펌용)').runs[0].font.size = Pt(9)

    for n in range(1, 29):
        e, k = en_p[n], ko_p[n]
        max_v = mt_verse_counts[n - 1]
        sections = parse_msg_sections(msg[str(n)], n, max_v)
        doc.add_heading(f'Matthew {n} — {e["title"]}', level=2)

        en_sets = [range_to_set(b) for b in e['verseRanges']]
        ko_sets = [range_to_set(b) for b in k['verseRanges']]
        en_paras = list(zip(e['paragraphs'], en_sets))
        ko_paras = list(zip(k['paragraphs'], ko_sets))

        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, t in enumerate(['절', 'The Message (원문)', 'Teenz EN (수정됨)', 'Teenz KO (수정됨)']):
            set_cell_text(hdr[i], t, bold=True, size=10, color=(255, 255, 255))
            shade(hdr[i], '1F3864')

        for rng, header, mtxt in sections:
            mset = range_to_set(rng)
            en_hit = [(idx, t) for idx, (t, s) in enumerate(en_paras) if s & mset]
            ko_hit = [(idx, t) for idx, (t, s) in enumerate(ko_paras) if s & mset]
            row = table.add_row().cells
            set_cell_text(row[0], rng, bold=True, size=9)
            first = True
            if header:
                add_para(row[1], header, bold=True, size=9, first=True)
                first = False
            add_para(row[1], mtxt, size=9, first=first)
            # EN cell
            if en_hit:
                f = True
                for idx, t in en_hit:
                    is_hdr = t.startswith('§')
                    is_new = norm(t) not in old_en_sets[n]
                    label = ('✚ ' if (is_new and not is_hdr) else '') + (t[1:] if is_hdr else t)
                    add_para(row[2], label, bold=is_hdr, size=9, first=f)
                    f = False
                if any(norm(t) not in old_en_sets[n] for _, t in en_hit):
                    shade(row[2], 'E2EFDA')
            else:
                set_cell_text(row[2], '⚠ 없음', size=9, color=(156, 0, 0), italic=True)
                shade(row[2], 'FFC7CE')
            # KO cell
            if ko_hit:
                f = True
                for idx, t in ko_hit:
                    is_hdr = t.startswith('§')
                    is_new = norm(t) not in old_ko_sets[n]
                    label = ('✚ ' if (is_new and not is_hdr) else '') + (t[1:] if is_hdr else t)
                    add_para(row[3], label, bold=is_hdr, size=9, first=f)
                    f = False
                if any(norm(t) not in old_ko_sets[n] for _, t in ko_hit):
                    shade(row[3], 'E2EFDA')
            else:
                set_cell_text(row[3], '⚠ 없음', size=9, color=(156, 0, 0), italic=True)
                shade(row[3], 'FFC7CE')
        doc.add_paragraph('')

    # 부록 A: 수정 내역
    doc.add_heading('부록 A: 수정 내역 (장별)', level=1)
    doc.add_paragraph('각 장에서 이번 수정 작업으로 바뀐 내용을 정리했습니다.').runs[0].font.size = Pt(9)
    for n in range(1, 29):
        changes = en_p[n].get('changes', []) + [f'[KO] {c}' for c in ko_p[n].get('changes', [])]
        if not changes:
            continue
        doc.add_heading(f'Matthew {n}', level=3)
        for c in changes:
            doc.add_paragraph(c, style='List Bullet')

    # 부록 B: 컨펌 필요
    doc.add_heading('부록 B: 컨펌 필요 목록 (MSG 문단 합병)', level=1)
    doc.add_paragraph(
        '원칙 2에 따라, MSG의 여러 문단을 하나의 Teenz 문단으로 합친 경우는 아래에 따로 표시했습니다. '
        '유지할지 나눌지 성욱이 직접 컨펌해 주세요.'
    ).runs[0].font.size = Pt(9)
    t2 = doc.add_table(rows=1, cols=3)
    t2.style = 'Table Grid'
    for i, h in enumerate(['장', '합병된 MSG 구간', '사유']):
        set_cell_text(t2.rows[0].cells[i], h, bold=True, size=10, color=(255, 255, 255))
        shade(t2.rows[0].cells[i], '1F3864')
    any_merge = False
    for n in range(1, 29):
        for m in en_p[n].get('merges', []):
            any_merge = True
            row = t2.add_row().cells
            set_cell_text(row[0], f'Matthew {n}', size=9)
            set_cell_text(row[1], ', '.join(m.get('msg_ranges', [])), size=9)
            set_cell_text(row[2], m.get('note', ''), size=9)
    if not any_merge:
        doc.add_paragraph('합병 없음 — 전 장이 MSG 문단 구분을 그대로 따릅니다.')

    # 부록 C: 검증 게이트 결과
    doc.add_heading('부록 C: 자동 검증 게이트 결과', level=1)
    try:
        r = subprocess.run(
            ['python3', os.path.join(WORK, 'validate_translation.py')]
            + sorted(glob.glob(os.path.join(FIX, 'en_matthew_*.json')))
            + sorted(glob.glob(os.path.join(FIX, 'ko_matthew_*.json'))),
            capture_output=True, text=True, timeout=60)
        out = (r.stdout + r.stderr).strip() or '(출력 없음)'
    except Exception as ex:
        out = f'게이트 실행 실패: {ex}'
    doc.add_paragraph(out).runs[0].font.size = Pt(9)

    out_path = os.path.join(OUT_DIR, 'Teenz-Bible-마태복음-번역대조-수정본.docx')
    doc.save(out_path)
    print('saved:', out_path)


if __name__ == '__main__':
    main()
