#!/usr/bin/env python3
"""Re-upload rebuilt DOCX files into the EXISTING Google Doc IDs (links stay the same),
then verify content via export: table counts, no empty Teen cells in non-header rows,
and the no-drop gate (every Teen EN/KO paragraph + every MSG unit present).

Usage: python3 reupload.py Matthew Mark Luke John Romans
"""
import json, subprocess, sys, time, os

GDIR = '/home/hatch/workspace/teenz-bible-review/gdocs_build'
REVIEW = '/home/hatch/workspace/teenz-bible-review'
sys.path.insert(0, GDIR)
from build_gdocs import (load_teen, parse_msg_txt, parse_msg_matthew,
                         parse_msg_romans, verify_no_drop)

BOOK_TABLES = {
    'Matthew': 28, 'Mark': 16, 'Luke': 24, 'John': 21, 'Acts': 28, 'Romans': 16,
    '1Corinthians': 16, '2Corinthians': 13, 'Galatians': 6, 'Ephesians': 6,
    'Philippians': 4, 'Colossians': 4, '1Thessalonians': 5, '2Thessalonians': 3,
    '1Timothy': 6, '2Timothy': 4, 'Titus': 3, 'Philemon': 1, 'Hebrews': 13,
    'James': 5, '1Peter': 5, '2Peter': 3, '1John': 5, '2John': 1, '3John': 1,
    'Jude': 1, 'Revelation': 22,
}

def run(args):
    r = subprocess.run(args, capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(f'CMD FAILED: {" ".join(args[:6])}...\nSTDERR: {r.stderr[:600]}')
    out = r.stdout.strip()
    return json.loads(out) if out else {}

def get_units(key):
    if key == 'Matthew':
        return parse_msg_matthew()
    if key == 'Romans':
        return parse_msg_romans()
    return parse_msg_txt(os.path.join(REVIEW, f'msg_{key}.txt'))

def cell_text(cell):
    return '\n'.join(p.text for p in cell.paragraphs).strip()

def cell_is_bold(cell):
    for p in cell.paragraphs:
        for r in p.runs:
            if r.bold:
                return True
    return False

def verify_exported(key, path):
    """Check the docx exported back from Google Docs."""
    from docx import Document
    doc = Document(path)
    tables = doc.tables
    problems = []
    exp = BOOK_TABLES[key]
    if len(tables) != exp:
        problems.append(f'table count {len(tables)} != {exp}')
    for ti, t in enumerate(tables):
        rows = t.rows
        if not rows:
            problems.append(f'table {ti}: no rows')
            continue
        for ri, row in enumerate(rows[1:], start=1):  # skip header row
            c0, c1, c2 = (cell_text(c) for c in row.cells)
            if not c0:
                # must be a § header row: Teen cells carry the header
                if not c1 or not c2:
                    problems.append(f'table {ti} row {ri}: empty MSG cell but Teen cells also empty')
            else:
                if not c1:
                    problems.append(f'table {ti} row {ri}: MSG present but Teen EN empty')
                if not c2:
                    problems.append(f'table {ti} row {ri}: MSG present but Teen KO empty')
    # no-drop gate on the exported file
    en_chs = load_teen('en', key)
    ko_chs = load_teen('ko', key)
    missing = verify_no_drop(key, path, en_chs, ko_chs, get_units(key))
    for m in missing:
        problems.append(f'DROPPED {m[0]} ch{m[1]}: {m[2][:50]!r}')
    return problems

def main():
    keys = sys.argv[1:]
    assert keys, 'usage: reupload.py KEY...'
    results = json.load(open(os.path.join(GDIR, 'upload_results.json'), encoding='utf-8'))
    id_by_key = {r['key']: r['id'] for r in results if r.get('id')}
    summary = []
    for key in keys:
        doc_id = id_by_key.get(key)
        assert doc_id, f'no doc id for {key}'
        docx = os.path.join(GDIR, f'{key}.docx')
        assert os.path.exists(docx) and os.path.getsize(docx) > 0, f'missing {docx}'
        print(f'[{key}] uploading to {doc_id} ...', flush=True)
        run(['hatch_gws_cli', 'drive', 'files', 'update',
             '--params', json.dumps({'fileId': doc_id}), '--upload', docx])
        time.sleep(3)
        # quick API check: table count
        g = run(['hatch_gws_cli', 'docs', 'documents', 'get',
                 '--params', json.dumps({'documentId': doc_id,
                                         'fields': 'title,body(content(table))'})])
        n_tables = sum(1 for x in g.get('body', {}).get('content', []) if 'table' in x)
        exp = BOOK_TABLES[key]
        print(f'[{key}] API table count: {n_tables}/{exp}', flush=True)
        # export back and deep-verify
        exp_path = f'/tmp/gverify_{key}.docx'
        r = subprocess.run(['hatch_gws_cli', 'drive', 'files', 'export',
                            '--params', json.dumps({'fileId': doc_id,
                                                    'mimeType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}),
                            '--output', exp_path],
                           capture_output=True, text=True, timeout=300)
        if r.returncode != 0:
            print(f'[{key}] EXPORT FAILED: {r.stderr[:300]}', flush=True)
            summary.append((key, ['export failed']))
            continue
        problems = verify_exported(key, exp_path)
        if problems:
            print(f'[{key}] PROBLEMS ({len(problems)}):', flush=True)
            for p in problems[:12]:
                print(f'    - {p}', flush=True)
        else:
            print(f'[{key}] VERIFIED OK', flush=True)
        summary.append((key, problems))
        time.sleep(1)
    n_ok = sum(1 for _, p in summary if not p)
    print(f'\nDONE: {n_ok}/{len(summary)} fully verified')
    bad = [(k, p) for k, p in summary if p]
    if bad:
        sys.exit(1)

if __name__ == '__main__':
    main()
