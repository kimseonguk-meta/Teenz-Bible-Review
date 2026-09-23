#!/usr/bin/env python3
"""Content-level doc pairing verification for Teenz Bible gdocs.

For each chapter table in a built/exported docx:
  - recompute the expected MSG units per teen paragraph (same greedy
    matcher as build_gdocs.match_msg_for_teen)
  - check the MSG cell actually contains the expected MSG unit text
    (or a 'continued' reference when the unit is shared with the row above)
  - check the EN / KO cells contain the expected teen paragraph text
    (badge first paragraph stripped, whitespace-normalized fuzzy match)

This catches text-level mispairings (e.g. the Matthew 1 genealogy incident)
that verse-coverage checks cannot see.

Usage: python3 verify_pairing_content.py <docx_path> <BookKey>
Exit 0 = all rows verified, 1 = problems found.
"""
import json, os, re, sys

GDIR = os.path.dirname(os.path.abspath(__file__))
REVIEW = os.path.dirname(GDIR)
sys.path.insert(0, GDIR)
import build_gdocs as bg
from docx import Document


def norm(t):
    return re.sub(r'\s+', ' ', (t or '')).strip()


def cell_paras(cell):
    return [p.text for p in cell.paragraphs]


def body_text(cell):
    """Cell text minus the badge first paragraph."""
    ps = cell_paras(cell)
    if ps and re.fullmatch(r'\d{1,2}(?:-\d{1,2})?(?:,\s*\d{1,2}(?:-\d{1,2})?)*', ps[0].strip()):
        ps = ps[1:]
    return '\n'.join(ps)


def main():
    docx_path, key = sys.argv[1], sys.argv[2]
    en_chs = bg.load_teen('en', key)
    ko_chs = bg.load_teen('ko', key)
    units = bg.parse_msg_txt(os.path.join(REVIEW, f'msg_{key}.txt'))
    by_ch = {}
    for u in units:
        by_ch.setdefault(u['chapter'], []).append(u)

    doc = Document(docx_path)
    problems = []
    checked = 0
    for ti, table in enumerate(doc.tables):
        chn = en_chs[ti]['chapter'] if ti < len(en_chs) else None
        ce, ck = en_chs[ti], ko_chs[ti]
        n = len(ce['paragraphs'])
        teen_infos = []
        for idx in range(n):
            pe = ce['paragraphs'][idx]
            vr = ce['verseRanges'][idx] if idx < len(ce['verseRanges']) else None
            teen_infos.append({'idx': idx,
                               'vset': bg.parse_verse_set(vr) if vr else set(),
                               'is_header': pe.startswith('§'), 'vr': vr})
        match, orphans, twm = bg.match_msg_for_teen(by_ch.get(chn, []), teen_infos)
        # expected MSG text per teen idx (None for headers)
        exp_msg = {}
        for idx in range(n):
            if teen_infos[idx]['is_header']:
                exp_msg[idx] = None
            else:
                exp_msg[idx] = '\n\n'.join('\n'.join(u['blocks']) for u in match[idx])

        rows = table.rows
        # row 0 = header; data rows follow in teen paragraph order,
        # plus possible orphan rows at the end
        data_row = 1
        for idx in range(n):
            if data_row >= len(rows):
                problems.append(f'ch{chn}: missing row for teen idx{idx}')
                break
            cells = rows[data_row].cells
            c0, c1, c2 = (norm(body_text(c)) for c in cells)
            pe = norm(ce['paragraphs'][idx].lstrip('§').strip())
            pk = norm(ck['paragraphs'][idx].lstrip('§').strip())
            if teen_infos[idx]['is_header']:
                if c0:
                    problems.append(f'ch{chn} idx{idx}: header row has MSG text')
                if pe[:60] not in c1:
                    problems.append(f'ch{chn} idx{idx}: header EN text mismatch')
                if pk[:60] not in c2:
                    problems.append(f'ch{chn} idx{idx}: header KO text mismatch')
            else:
                em = norm(exp_msg[idx])
                if not c0:
                    # Independent gate: a non-header teen row must always
                    # show MSG text. A badge-only MSG cell is a defect even
                    # when the parser itself produced no expected unit
                    # (that parser gap is what let this defect pass before).
                    problems.append(
                        f'ch{chn} idx{idx}: empty MSG cell (badge-only). '
                        f'teen vr={teen_infos[idx]["vr"]!r} '
                        f'parser_unit_empty={not em}')
                elif not em:
                    problems.append(
                        f'ch{chn} idx{idx}: MSG cell text not accounted for by parser '
                        f'(cell head={c0[:60]!r}) — pairing unverified')
                else:
                    head = em[:60]
                    if head not in c0 and not c0.startswith('↑ 위 MSG'):
                        problems.append(
                            f'ch{chn} idx{idx}: MSG cell does not contain expected unit text. '
                            f'expected head={head!r} got={c0[:60]!r}')
                    if c0.startswith('↑ 위 MSG'):
                        # continued reference: walk back through chained
                        # continued-ref rows to the row holding the full block
                        pr = data_row - 1
                        prev_c0 = norm(body_text(rows[pr].cells[0]))
                        while prev_c0.startswith('↑ 위 MSG') and pr > 1:
                            pr -= 1
                            prev_c0 = norm(body_text(rows[pr].cells[0]))
                        if head not in prev_c0:
                            problems.append(
                                f'ch{chn} idx{idx}: continued-ref but prev row lacks the block')
                if pe[:80] not in c1 and c1[:80] not in pe:
                    problems.append(f'ch{chn} idx{idx}: EN cell text mismatch')
                if pk[:80] not in c2 and c2[:80] not in pk:
                    problems.append(f'ch{chn} idx{idx}: KO cell text mismatch')
                checked += 1
            data_row += 1
        # orphan rows (if any) at end
        for oi, u in enumerate(orphans):
            if data_row + oi >= len(rows):
                problems.append(f'ch{chn}: missing orphan row {oi}')
    print(f'checked {checked} data rows in {len(doc.tables)} tables')
    if problems:
        print(f'PROBLEMS ({len(problems)}):')
        for p in problems[:20]:
            print(' -', p)
        sys.exit(1)
    print('PAIRING CONTENT OK')


if __name__ == '__main__':
    main()
