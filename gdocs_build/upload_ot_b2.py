#!/usr/bin/env python3
"""Worker B2: OT 4권(Psalms/Proverbs/Ecclesiastes/SongOfSongs) Google Docs 업로드.

Usage: python3 upload_ot_b2.py SongOfSongs [Psalms ...]

게이트 7: 폴더 1I7rLOsUgbr5Ewyq7-IEshCk20scIBtPn에
'{한글이름} ({영어이름}) — MSG + Teen EN + KO' 제목으로 Doc 생성
-> docx 업로드(자동 변환) -> Docs API title 일치 + 테이블 수 == 장 수 확인
-> export-back 검증 (빈 틴 셀 없음, 누락 문단 없음).
결과는 gdocs_build/upload_results_workerB2.json에 추가(키 단위 upsert).
(공유 upload_results.json은 동결 — 절대 쓰지 않음.)
"""
import json, subprocess, sys, time, os

GDIR = '/home/hatch/workspace/teenz-bible-review/gdocs_build'
REVIEW = '/home/hatch/workspace/teenz-bible-review'
sys.path.insert(0, GDIR)
from build_gdocs import load_teen, parse_msg_txt, verify_no_drop
from docx import Document

FOLDER = '1I7rLOsUgbr5Ewyq7-IEshCk20scIBtPn'

BOOKS_B2 = [
    ('Psalms', '시편', 'Psalms', 150),
    ('Proverbs', '잠언', 'Proverbs', 31),
    ('Ecclesiastes', '전도서', 'Ecclesiastes', 12),
    ('SongOfSongs', '아가', 'Song of Songs', 8),
]
BOOK_MAP = {k: (ko, en, n) for k, ko, en, n in BOOKS_B2}

def verify_exported_ot(key, path):
    doc = Document(path)
    tables = doc.tables
    problems = []
    exp = BOOK_MAP[key][2]
    if len(tables) != exp:
        problems.append(f'table count {len(tables)} != {exp}')
    def cell_text(cell):
        return '\n'.join(p.text for p in cell.paragraphs).strip()
    for ti, t in enumerate(tables):
        rows = t.rows
        if not rows:
            problems.append(f'table {ti}: no rows'); continue
        for ri, row in enumerate(rows[1:], start=1):
            c0, c1, c2 = (cell_text(c) for c in row.cells)
            if not c0:
                if not c1 or not c2:
                    problems.append(f'table {ti} row {ri}: empty MSG cell but Teen cells also empty')
            else:
                if not c1: problems.append(f'table {ti} row {ri}: MSG present but Teen EN empty')
                if not c2: problems.append(f'table {ti} row {ri}: MSG present but Teen KO empty')
    en_chs = load_teen('en', key); ko_chs = load_teen('ko', key)
    units = parse_msg_txt(os.path.join(REVIEW, f'msg_{key}.txt'))
    for m in verify_no_drop(key, path, en_chs, ko_chs, units):
        problems.append(f'DROPPED {m[0]} ch{m[1]}: {m[2][:50]!r}')
    return problems

def run(args):
    r = subprocess.run(args, capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(f'CMD FAILED: {" ".join(args[:6])}...\nSTDERR: {r.stderr[:600]}')
    out = r.stdout.strip()
    return json.loads(out) if out else {}

def main():
    keys = sys.argv[1:]
    assert keys and all(k in BOOK_MAP for k in keys), 'usage: upload_ot_b2.py <OT-B2 Key...>'
    res_path = os.path.join(GDIR, 'upload_results_workerB2.json')
    results = json.load(open(res_path, encoding='utf-8')) if os.path.exists(res_path) else []
    id_by_key = {r['key']: r for r in results}
    summary = []
    for key in keys:
        ko, en, exp_tables = BOOK_MAP[key]
        title = f'{ko} ({en}) — MSG + Teen EN + KO'
        docx = os.path.join(GDIR, f'{key}.docx')
        assert os.path.exists(docx) and os.path.getsize(docx) > 0, f'missing {docx}'
        print(f'[{key}] creating doc...', flush=True)
        try:
            c = run(['hatch_gws_cli', 'drive', 'files', 'create',
                     '--params', '{"ignoreDefaultVisibility":true}',
                     '--json', json.dumps({'name': title,
                                           'mimeType': 'application/vnd.google-apps.document',
                                           'parents': [FOLDER]})])
            doc_id = c['id']
            time.sleep(1)
            run(['hatch_gws_cli', 'drive', 'files', 'update',
                 '--params', json.dumps({'fileId': doc_id}), '--upload', docx])
            time.sleep(3)
            g = run(['hatch_gws_cli', 'docs', 'documents', 'get',
                     '--params', json.dumps({'documentId': doc_id,
                                             'fields': 'title,body(content(table))'})])
            n_tables = sum(1 for x in g.get('body', {}).get('content', []) if 'table' in x)
            title_ok = (g.get('title') == title)
            print(f'[{key}] API: title_ok={title_ok} tables={n_tables}/{exp_tables}', flush=True)
            exp_path = f'/tmp/gverify_{key}.docx'
            r = subprocess.run(['hatch_gws_cli', 'drive', 'files', 'export',
                                '--params', json.dumps({'fileId': doc_id,
                                    'mimeType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}),
                                '--output', exp_path],
                               capture_output=True, text=True, timeout=300)
            problems = []
            if r.returncode != 0:
                problems.append('export failed: ' + r.stderr[:200])
            else:
                problems = verify_exported_ot(key, exp_path)
            f = run(['hatch_gws_cli', 'drive', 'files', 'get',
                     '--params', json.dumps({'fileId': doc_id,
                                             'fields': 'id,name,webViewLink,parents'})])
            ok = title_ok and n_tables == exp_tables and not problems
            entry = {'key': key, 'title': title, 'id': doc_id,
                     'webViewLink': f.get('webViewLink'),
                     'tables': n_tables, 'expected': exp_tables,
                     'status': 'OK' if ok else 'MISMATCH',
                     'export_problems': problems[:10],
                     'export_back': 'VERIFIED OK' if (ok and not problems) else 'ISSUES'}
            if key in id_by_key:
                results[results.index(id_by_key[key])] = entry
            else:
                results.append(entry)
            id_by_key[key] = entry
            print(f'[{key}] -> {entry["status"]}' + ('' if ok else f' {problems[:5]}'), flush=True)
            summary.append((key, ok))
        except Exception as e:
            print(f'[{key}] FAILED: {str(e)[:300]}', flush=True)
            summary.append((key, False))
        time.sleep(1)
    json.dump(results, open(res_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    n_ok = sum(1 for _, ok in summary if ok)
    print(f'\nDONE: {n_ok}/{len(summary)} OK')
    if n_ok != len(summary):
        sys.exit(1)

if __name__ == '__main__':
    main()
