#!/usr/bin/env python3
"""Build one .docx per NT book: 3-column comparison MSG | Teen EN | Teen KO.

Read-only wrt source patch files and MSG txt files. Output -> gdocs_build/<Key>.docx
Usage: build_gdocs.py [Key ...]   (default: all 27 books)

MSG txt formats observed (BibleGateway fetches):
  Format A: '# ...' + 'Source: ...' headers, '### ' subtitles, chapter markers
            'N M-...' (two numbers), verse paragraphs 'M-...'.
  Format B: '<Book> <range> MSG -- source: ...' header line, lone-number
            chapter markers on their own line ('1', '2', ...), verse
            paragraphs 'M-...', '### ' subtitles inline.
  Mixed:    single-line chapter markers 'N <text>' (chapter N, verse 1),
            e.g. Romans '1 I, Paul, ...', Philippians '3 And that's about it'.
Verse-less continuation blocks (psalm quotes, poetry, genealogy lines)
are attached to the current MSG unit so no text is lost. '* * *' and
'### ' lines are dropped.
"""
import json, re, os, sys
from docx import Document
from docx.shared import Pt, Inches

REVIEW = '/home/hatch/workspace/teenz-bible-review'
BUILD = os.path.join(REVIEW, 'gdocs_build')
os.makedirs(BUILD, exist_ok=True)

BOOKS = [
    ('Matthew', '마태복음', 'Matthew'),
    ('Mark', '마가복음', 'Mark'),
    ('Luke', '누가복음', 'Luke'),
    ('John', '요한복음', 'John'),
    ('Acts', '사도행전', 'Acts'),
    ('Romans', '로마서', 'Romans'),
    ('1Corinthians', '고린도전서', '1 Corinthians'),
    ('2Corinthians', '고린도후서', '2 Corinthians'),
    ('Galatians', '갈라디아서', 'Galatians'),
    ('Ephesians', '에베소서', 'Ephesians'),
    ('Philippians', '빌립보서', 'Philippians'),
    ('Colossians', '골로새서', 'Colossians'),
    ('1Thessalonians', '데살로니가전서', '1 Thessalonians'),
    ('2Thessalonians', '데살로니가후서', '2 Thessalonians'),
    ('1Timothy', '디모데전서', '1 Timothy'),
    ('2Timothy', '디모데후서', '2 Timothy'),
    ('Titus', '디도서', 'Titus'),
    ('Philemon', '빌레몬서', 'Philemon'),
    ('Hebrews', '히브리서', 'Hebrews'),
    ('James', '야고보서', 'James'),
    ('1Peter', '베드로전서', '1 Peter'),
    ('2Peter', '베드로후서', '2 Peter'),
    ('1John', '요한일서', '1 John'),
    ('2John', '요한이서', '2 John'),
    ('3John', '요한삼서', '3 John'),
    ('Jude', '유다서', 'Jude'),
    ('Revelation', '요한계시록', 'Revelation'),
]
BOOK_MAP = {k: (ko, en) for k, ko, en in BOOKS}

TEXT_START = r'[A-Z"“‘(\[*]'
VERSE_ATOM = r'\d{1,2}(?:-\d{1,2})?'
VERSE_REF = VERSE_ATOM + r'(?:,\s*' + VERSE_ATOM + r')*'
VERSE_RE = re.compile(r'^(' + VERSE_REF + r')\s+(?=' + TEXT_START + r')(.*)$')
TWO_NUM_RE = re.compile(r'^(\d{1,3})\s+(' + VERSE_REF + r')\s+(?=' +
                        TEXT_START + r')(.*)$')
LONE_NUM_RE = re.compile(r'^(\d{1,3})$')
SINGLE_NUM_TEXT_RE = re.compile(r'^(\d{1,3})\s+(?=' + TEXT_START + r')(.*)$')
# Inline verse-range markers embedded mid-line, e.g.
#   "1 The family tree of Jesus Christ, David's son, Abraham's son: 2-6 Abraham had Isaac,"
#   "Soon the whole world will be taunting them: 6-8 "'Who do you think you are-"
# The (?<!\d) guard keeps chapter:verse references like 'Mark 16:9-20'
# (a textual note, not a marker) untouched.
INLINE_MARKER_RE = re.compile(r'(?<!\d):\s*(\d{1,2}-\d{1,2})\s+(?=[A-Z"“‘(\[])')


def expand_inline_markers(text):
    """Turn inline ': N-M Text' markers into a line break so the normal
    VERSE classifier starts a new MSG unit there.

    Without this, the marker's text is glued to the previous unit's tail and
    the doc pairs it with the wrong Teen paragraph (e.g. Matthew 1's
    genealogy displayed next to the chapter summary, making the correct
    translation look like an arbitrary abbreviation). Header/source lines are
    left untouched.
    """
    out_lines = []
    for line in text.split('\n'):
        s = line.strip()
        low = s.lower()
        if (s.startswith('#') or s.startswith('Source:') or s == '* * *'
                or ('msg' in low and 'source:' in low)):
            out_lines.append(line)
            continue
        out_lines.append(INLINE_MARKER_RE.sub(r':\n\1 ', line))
    return '\n'.join(out_lines)


def parse_verse_set(s):
    if not s:
        return set()
    s = s.replace('–', '-').replace('—', '-')
    out = set()
    for part in re.split(r'[,\s]+', s.strip()):
        part = part.strip().strip(',;')
        if not part:
            continue
        m = re.match(r'^(\d+)-(\d+)$', part)
        if m:
            a, b = int(m.group(1)), int(m.group(2))
            if a <= b:
                out.update(range(a, b + 1))
        elif part.isdigit():
            out.add(int(part))
    return out


def classify_lines(text):
    """Pass 1: classify each line. Returns list of (kind, payload)."""
    out = []
    for raw in text.split('\n'):
        line = raw.strip()
        if not line:
            out.append(('BLANK', None))
            continue
        low = line.lower()
        if (line.startswith('#') or line.startswith('Source:') or
                line == '* * *' or
                ('msg' in low and 'source:' in low)):
            out.append(('SKIP', None))
            continue
        m = LONE_NUM_RE.match(line)
        if m:
            out.append(('LONE_NUM', int(m.group(1))))
            continue
        m = TWO_NUM_RE.match(line)
        if m:
            out.append(('TWO_NUM', (int(m.group(1)), m.group(2), m.group(3))))
            continue
        m = VERSE_RE.match(line)
        if m:
            out.append(('VERSE', (m.group(1), m.group(2))))
            continue
        m = SINGLE_NUM_TEXT_RE.match(line)
        if m:
            out.append(('SINGLE_NUM_TEXT', (int(m.group(1)), m.group(2))))
            continue
        out.append(('TEXT', line))
    return out


def next_verse_mins(classified):
    """For each line index, min verse of the *immediately next* VERSE/TWO_NUM
    line (closest following one, not the global minimum)."""
    n = len(classified)
    res = [None] * n
    best = None
    for i in range(n - 1, -1, -1):
        res[i] = best
        kind, pay = classified[i]
        if kind == 'VERSE':
            v = parse_verse_set(pay[0])
            best = min(v) if v else None
        elif kind == 'TWO_NUM':
            v = parse_verse_set(pay[1])
            best = min(v) if v else None
    return res


def parse_msg_lines(text, initial_ch=0):
    """Pass 2: build MSG units [{chapter, verses:set, blocks:[str]}]."""
    text = expand_inline_markers(text)
    classified = classify_lines(text)
    nxt = next_verse_mins(classified)
    units = []
    cur_ch = initial_ch
    covered = set()
    cur_maxv = 0
    cur_unit = None
    at_block_start = True

    def new_unit(ch, verses, first_block):
        nonlocal cur_unit, covered, cur_maxv
        cur_unit = {'chapter': ch, 'verses': set(verses),
                    'blocks': [first_block] if first_block else []}
        units.append(cur_unit)
        covered |= set(verses)
        if verses:
            cur_maxv = max(cur_maxv, max(verses))

    def chapter_marker(ch, rest_text):
        nonlocal cur_ch, covered, cur_maxv
        cur_ch = ch
        covered = set()
        cur_maxv = 0
        # single-line markers ('N <text>') print verse 1 without its number
        new_unit(ch, {1}, rest_text)

    for i, (kind, pay) in enumerate(classified):
        if kind in ('BLANK', 'SKIP'):
            if kind == 'BLANK':
                at_block_start = True
            continue
        if kind == 'LONE_NUM':
            n = pay
            if n > cur_ch:
                cur_ch = n
                covered = set()
                cur_maxv = 0
                cur_unit = None
            # else: stray number line, ignore
            continue
        if kind == 'TWO_NUM':
            c, vstr, rest = pay
            if c > cur_ch:
                cur_ch = c
                covered = set()
                cur_maxv = 0
            new_unit(cur_ch, parse_verse_set(vstr), rest.strip())
            at_block_start = False
            continue
        if kind in ('VERSE', 'SINGLE_NUM_TEXT'):
            if kind == 'VERSE':
                vstr, rest = pay
            else:
                n, rest = pay
                vstr = str(n)
            vset = parse_verse_set(vstr)
            is_single_num = (len(vset) == 1 and
                             re.fullmatch(r'\d+', vstr.strip()) is not None)
            if cur_ch == 0:
                if is_single_num and int(vstr) >= 1:
                    # file must open with a chapter, e.g. Romans '1 I, Paul...'
                    chapter_marker(int(vstr), rest.strip())
                    at_block_start = False
                # else: orphan verse line before any chapter -> drop
                continue
            if is_single_num:
                nn = int(vstr)
                if nn == cur_ch and not covered:
                    # Chapter opens with 'N <text>' where N is the chapter
                    # number: MSG prints verse 1 without its own number
                    # (e.g. Matthew 6/11/18: '6 "Be especially careful...").
                    # Without this, verse-1 text is misfiled as verse N.
                    new_unit(cur_ch, {1}, rest.strip())
                    at_block_start = False
                    continue
                if nn == cur_ch + 1:
                    # ambiguous: verse nn of cur_ch, or chapter nn marker?
                    if nn in covered:
                        chapter_marker(nn, rest.strip())
                    elif nxt[i] is None:
                        # last verse line of the file: a chapter marker
                        # would have following content, so this is a verse
                        # (e.g. 2Cor 13:14 printed as '14 ...')
                        new_unit(cur_ch, vset, rest.strip())
                    elif nxt[i] > nn:
                        # verses continue forward -> verse nn of cur_ch
                        # (e.g. Acts 6:7 '...', next is '8-10 ...')
                        new_unit(cur_ch, vset, rest.strip())
                    else:
                        # verses restart low -> new chapter
                        # (e.g. Acts 7:1 '...', next is '2-3 ...')
                        chapter_marker(nn, rest.strip())
                    at_block_start = False
                    continue
            new_unit(cur_ch, vset, rest.strip())
            at_block_start = False
            continue
        # TEXT
        if cur_unit is None:
            continue  # orphan text before any verse -> drop
        if at_block_start:
            cur_unit['blocks'].append(pay)
        else:
            cur_unit['blocks'][-1] = cur_unit['blocks'][-1] + ' ' + pay
        at_block_start = False
    return units


def parse_msg_txt(path):
    return parse_msg_lines(open(path, encoding='utf-8').read())


def parse_msg_matthew():
    chunks = []
    raw = json.load(open(os.path.join(REVIEW, 'msg_matthew_raw.json'),
                         encoding='utf-8'))
    for ch_str in sorted(raw.keys(), key=int):
        chunks.append((int(ch_str), raw[ch_str]))
    for ch, fname in ((1, 'msg_raw_1.txt'), (3, 'msg_raw_3.txt')):
        chunks.append((ch, open(os.path.join(REVIEW, fname),
                                encoding='utf-8').read()))
    units = []
    for ch, text in chunks:
        units.extend(parse_msg_lines(text, initial_ch=ch))
    units.sort(key=lambda u: (u['chapter'],))
    return units


def load_teen(lang, key):
    fx = os.path.join(REVIEW, 'fixes')
    if key == 'Matthew':
        if lang == 'en':
            files = ['en_matthew_01-03.json', 'en_matthew_04-07.json',
                     'en_matthew_08-10.json', 'en_matthew_11-14.json',
                     'en_matthew_15-28.json']
        else:
            files = ['ko_matthew_01-14.json', 'ko_matthew_15-28.json']
        chs = []
        for f in files:
            chs.extend(json.load(open(os.path.join(fx, f), encoding='utf-8')))
        return chs
    return json.load(open(os.path.join(fx, f'{lang}_{key}.json'),
                          encoding='utf-8'))


def match_msg_for_teen(ch_units, teen_infos):
    """Per-Teen-paragraph greedy-coverage match (sharing allowed).

    For each non-header Teen paragraph, take the MSG units that intersect
    its verse range, ordered by overlap (desc), and greedily keep a unit
    when it covers at least one verse not yet covered. This handles:
      - splits (several Teen paragraphs share one MSG unit, Principle 2),
      - merges (one Teen paragraph takes several MSG units),
      - MSG's overlapping printed ranges ('1-3' then '3-6' share v.3:
        the second unit adds no new verse, so no duplication).
    MSG units selected by no Teen paragraph are returned as orphans so the
    caller can surface them instead of silently dropping them (Principle 3).

    teen_infos: list of {'idx', 'vset' (set), 'is_header', 'vr'}.
    Returns (match_dict, orphan_units, teen_without_msg).
    """
    match = {ti['idx']: [] for ti in teen_infos}
    used = set()
    teen_without_msg = []
    for ti in teen_infos:
        if ti['is_header'] or not ti['vset']:
            continue
        tset = ti['vset']
        cands = []
        for ui, u in enumerate(ch_units):
            uverses = u['verses'] or set()
            ov = len(uverses & tset)
            if ov > 0:
                # Tiebreak equal absolute overlap by coverage ratio: prefer the
                # unit whose verses sit most fully inside the Teen badge.
                # (Without this, a Teen badge that is a strict subset of a wider
                # MSG unit's range loses to that unit by index order, orphaning
                # the narrower exact-match unit — e.g. Galatians 2:21 vs 19-21.)
                cands.append((ui, ov, len(uverses)))
        cands.sort(key=lambda x: (-x[1], -x[1] / x[2], x[0]))
        covered, picked = set(), []
        for ui, _ov, _vl in cands:
            new = (ch_units[ui]['verses'] or set()) & tset - covered
            if new:
                picked.append(ui)
                covered |= (ch_units[ui]['verses'] or set()) & tset
        picked.sort()
        if picked:
            match[ti['idx']] = [ch_units[i] for i in picked]
            used.update(picked)
        else:
            teen_without_msg.append((ti['idx'], ti['vr']))
    orphans = [u for ui, u in enumerate(ch_units) if ui not in used]
    return match, orphans, teen_without_msg


def set_cell(cell, text, bold=False, size=Pt(9), badge=None):
    """Write text into a table cell.

    badge: verse-range string (e.g. '19-21') rendered as its own small bold
    first paragraph so the verse badge is actually visible in the document.
    Previously verseRanges were used only for MSG↔EN↔KO matching and never
    printed, so badges were invisible in Google Docs (structural builder bug).
    """
    cell.text = ''
    first = True
    blocks = []
    if badge:
        blocks.append(('__badge__', badge))
    for para_text in (text.split('\n') if text else ['']):
        blocks.append(('body', para_text))
    for kind, para_text in blocks:
        p = cell.paragraphs[0] if first else cell.add_paragraph()
        first = False
        run = p.add_run(para_text)
        if kind == '__badge__':
            run.bold = True
            run.font.size = Pt(8)
        else:
            run.font.size = size
            if bold:
                run.bold = True


def parse_msg_romans():
    """Romans: msg_Romans.txt is truncated mid-14:1; msg_Romans_supplement.txt
    (separate file; original untouched) completes 14:1-16:27."""
    main_units = [u for u in parse_msg_txt(os.path.join(REVIEW, 'msg_Romans.txt'))
                  if u['chapter'] != 14]
    sup_units = parse_msg_txt(os.path.join(REVIEW, 'msg_Romans_supplement.txt'))
    return main_units + sup_units


def mark_shared_msg_cells(table, data_rows):
    """Mark (not merge) the MSG cells of consecutive Teen rows that share
    the identical MSG unit list (e.g. three Teen paragraphs splitting one
    MSG 7-13 block). The first row keeps the full MSG block; the following
    rows show a short "continued" reference instead of repeating the block.

    NOTE: vertical cell merging was tried first but this environment's
    python-docx writes vMerge=restart on every cell of the span (0 continue
    cells) and concatenates the text, so merging is broken here. The
    reference-text approach renders identically in Word and Google Docs
    with no XML tricks. Without this, every row repeats the full MSG block
    and each Teen paragraph looks like a summary of the whole block.

    data_rows: list of (row_idx, teen_idx, unit_key, units, vr) for
    non-header rows, in document order.
    Returns groups (list of lists of data_rows entries) for pairing
    verification.
    """
    groups = []
    i, n = 0, len(data_rows)
    while i < n:
        j = i
        while (j + 1 < n
               and data_rows[j + 1][2]
               and data_rows[j + 1][2] == data_rows[i][2]
               and data_rows[j + 1][0] == data_rows[j][0] + 1):
            j += 1
        group = data_rows[i:j + 1]
        groups.append(group)
        if len(group) > 1:
            msg_v = set()
            for u in group[0][3]:
                msg_v |= set(u['verses'] or set())
            ref = '↑ 위 MSG 블록 계속'
            if msg_v:
                ref = f'↑ 위 MSG({min(msg_v)}-{max(msg_v)}) 계속'
            for (rr, _ti, _uk, _u, _vr) in group[1:]:
                set_cell(table.cell(rr, 0), ref, badge=_vr)
        i = j + 1
    return groups


def check_pairing(key, chn, groups):
    """Verify each displayed MSG cell's verse coverage matches the Teen
    paragraph badge(s) shown against it. Returns a list of mismatch dicts;
    mismatches are reported, never silently passed. Rows with no MSG match
    are skipped here (already reported as teen_without_msg).
    """
    mism = []
    for g in groups:
        units = g[0][3]
        if not units:
            continue
        msg_v = set()
        for u in units:
            msg_v |= set(u['verses'] or set())
        teen_v = set()
        vrs = [vr for (_r, _ti, _uk, _u, vr) in g]
        for vr in vrs:
            if vr:
                teen_v |= parse_verse_set(vr)
        if msg_v != teen_v:
            if teen_v and teen_v < msg_v:
                cat = 'MSG_SUPERSET'
            elif msg_v and msg_v < teen_v:
                cat = 'MSG_SUBSET'
            else:
                cat = 'MISMATCH'
            mism.append({'chapter': chn,
                         'teen_idx': [t[1] for t in g],
                         'badges': vrs,
                         'msg_verses': sorted(msg_v),
                         'teen_verses': sorted(teen_v),
                         'category': cat})
    return mism


def build_docx(key):
    ko_name, en_name = BOOK_MAP[key]
    en_chs = load_teen('en', key)
    ko_chs = load_teen('ko', key)
    assert len(en_chs) == len(ko_chs), f'{key}: chapter count mismatch'
    for ce, ck in zip(en_chs, ko_chs):
        assert ce['chapter'] == ck['chapter'], \
            f'{key}: chapter order mismatch'
        assert len(ce['paragraphs']) == len(ck['paragraphs']), \
            f'{key} ch{ce["chapter"]}: EN/KO paragraph count mismatch'

    msg_units = (parse_msg_matthew() if key == 'Matthew'
                 else parse_msg_romans() if key == 'Romans'
                 else parse_msg_txt(os.path.join(REVIEW, f'msg_{key}.txt')))
    by_ch = {}
    for u in msg_units:
        by_ch.setdefault(u['chapter'], []).append(u)

    doc = Document()
    doc.styles['Normal'].font.size = Pt(9.5)
    doc.add_heading(f'{ko_name} ({en_name})', level=1)
    doc.add_paragraph(
        'The Message 영어 원문 | Teen English 최종본 | Teen 한국어 최종본 — '
        '장별 3단 비교')

    total_rows = 0
    teen_without_msg = []
    msg_orphans = []
    pairing_mismatches = []
    for ce, ck in zip(en_chs, ko_chs):
        chn = ce['chapter']
        doc.add_heading(f'Chapter {chn} — {ce.get("title", "")}', level=1)
        n = len(ce['paragraphs'])
        vr_list = ce.get('verseRanges') or [None] * n
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Table Grid'
        hdr = table.rows[0].cells
        for i, t in enumerate(['MSG 영어 원문', 'Teen English', 'Teen 한국어']):
            set_cell(hdr[i], t, bold=True, size=Pt(10))
        for row in table.rows:
            for i in range(3):
                row.cells[i].width = Inches(2.17)
        teen_infos = []
        for idx in range(n):
            pe = ce['paragraphs'][idx]
            vr = vr_list[idx] if idx < len(vr_list) else None
            teen_infos.append({
                'idx': idx,
                'vset': parse_verse_set(vr) if vr else set(),
                'is_header': pe.startswith('§'),
                'vr': vr,
            })
        match, orphans, twm = match_msg_for_teen(by_ch.get(chn, []),
                                                  teen_infos)
        teen_without_msg.extend((chn, i, vr) for i, vr in twm)
        data_rows = []  # (row_idx, teen_idx, unit_key, units, vr)
        for idx in range(n):
            pe, pk = ce['paragraphs'][idx], ck['paragraphs'][idx]
            ti = teen_infos[idx]
            row = table.add_row().cells
            ridx = len(table.rows) - 1
            if ti['is_header']:
                # § section header: not MSG content; Teen carries the header.
                # Principle 2: the header shows the same badge as the
                # following body paragraph.
                hdr_badge = None
                for j in range(idx + 1, n):
                    if not teen_infos[j]['is_header']:
                        hdr_badge = teen_infos[j]['vr']
                        break
                set_cell(row[0], '')
                set_cell(row[1], pe.lstrip('§').strip(), bold=True,
                         badge=hdr_badge)
                set_cell(row[2], pk.lstrip('§').strip(), bold=True,
                         badge=hdr_badge)
            else:
                units = match[idx]
                msg_text = '\n\n'.join('\n'.join(u['blocks'])
                                       for u in units)
                set_cell(row[0], msg_text, badge=ti['vr'])
                set_cell(row[1], pe, badge=ti['vr'])
                set_cell(row[2], pk, badge=ti['vr'])
                data_rows.append((ridx, idx, tuple(id(u) for u in units),
                                  units, ti['vr']))
            total_rows += 1
        # Principle 3 safety net: never silently drop MSG text. Any MSG unit
        # that no Teen paragraph selected gets its own labeled row at the end
        # of the chapter table.
        for u in orphans:
            msg_orphans.append((chn, sorted(u['verses'])))
            row = table.add_row().cells
            uv = sorted(u['verses'] or set())
            orphan_badge = (f'{uv[0]}-{uv[-1]}' if len(uv) > 1
                            else (str(uv[0]) if uv else None))
            set_cell(row[0], '\n'.join(u['blocks']), badge=orphan_badge)
            set_cell(row[1], '(확인 필요: 대응 틴즈 문단 없음)', bold=True)
            set_cell(row[2], '(확인 필요: 대응 틴즈 문단 없음)', bold=True)
            total_rows += 1
        # Shared MSG units across consecutive Teen rows: show the full MSG
        # block once (first row) and a "continued" reference in the rest,
        # instead of repeating the block in every row.
        groups = mark_shared_msg_cells(table, data_rows)
        pairing_mismatches.extend(check_pairing(key, chn, groups))

    out = os.path.join(BUILD, f'{key}.docx')
    doc.save(out)
    return {'key': key, 'file': out, 'chapters': len(en_chs),
            'msg_chapters': sorted(by_ch.keys()),
            'rows': total_rows, 'teen_without_msg': teen_without_msg,
            'msg_orphans': msg_orphans,
            'msg_units': len(msg_units),
            'pairing_mismatches': pairing_mismatches}


def verify_no_drop(key, docx_path, en_chs, ko_chs, msg_units):
    """Principle 3 gate: every Teen EN paragraph, every Teen KO paragraph,
    and every MSG unit's text must appear verbatim in the built docx.
    Returns a list of ('EN'|'KO'|'MSG', chapter, snippet) for anything missing.
    """
    doc = Document(docx_path)
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                parts.append('\n'.join(p.text for p in cell.paragraphs))
    blob = ' '.join(' '.join(parts).split())

    def present(text):
        t = ' '.join(text.split())
        if not t:
            return True
        probe = t[:80]
        return probe in blob

    missing = []
    for ch in en_chs:
        for p in ch['paragraphs']:
            t = p.lstrip('§').strip() if p.startswith('§') else p.strip()
            if not present(t):
                missing.append(('EN', ch['chapter'], t[:60]))
    for ch in ko_chs:
        for p in ch['paragraphs']:
            t = p.lstrip('§').strip() if p.startswith('§') else p.strip()
            if not present(t):
                missing.append(('KO', ch['chapter'], t[:60]))
    for u in msg_units:
        t = ' '.join(u['blocks']).strip()
        if not present(t):
            missing.append(('MSG', u['chapter'],
                            (t[:60] if t else '<empty blocks>')))
    return missing


def main():
    keys = sys.argv[1:] or [k for k, _, _ in BOOKS]
    for k in keys:
        assert k in BOOK_MAP, f'unknown book key: {k}'
    all_missing = {}
    for k in keys:
        r = build_docx(k)
        en_chs = load_teen('en', k)
        ko_chs = load_teen('ko', k)
        msg_units = (parse_msg_matthew() if k == 'Matthew'
                     else parse_msg_romans() if k == 'Romans'
                     else parse_msg_txt(os.path.join(REVIEW, f'msg_{k}.txt')))
        missing = verify_no_drop(k, r['file'], en_chs, ko_chs, msg_units)
        all_missing[k] = missing
        exp = [c['chapter'] for c in en_chs]
        missing_ch = [c for c in exp if c not in r['msg_chapters']]
        print(f"{r['key']}: teen_ch={r['chapters']} msg_chapters={len(r['msg_chapters'])} "
              f"missing_msg_ch={missing_ch} rows={r['rows']} "
              f"msg_units={r['msg_units']} "
              f"teen_without_msg={len(r['teen_without_msg'])} "
              f"msg_orphans={len(r['msg_orphans'])} "
              f"VERIFY_DROPPED={len(missing)}")
        for m in missing[:10]:
            print(f"    DROPPED {m[0]} ch{m[1]}: {m[2]!r}")
        if r['teen_without_msg'][:5]:
            print(f"    teen_without_msg sample: {r['teen_without_msg'][:5]}")
        if r['msg_orphans'][:5]:
            print(f"    msg_orphans sample: {r['msg_orphans'][:5]}")
        pm = r['pairing_mismatches']
        if pm:
            print(f"    PAIRING_MISMATCHES={len(pm)} "
                  f"(MSG cell verse coverage != teen badge; review needed)")
            for m in pm[:15]:
                print(f"      ch{m['chapter']} teen_idx={m['teen_idx']} "
                      f"badges={m['badges']} msg_verses={m['msg_verses']} "
                      f"teen_verses={m['teen_verses']} [{m['category']}]")
    bad = {k: v for k, v in all_missing.items() if v}
    if bad:
        print(f"\nFAIL: content dropped in {sorted(bad)}")
        sys.exit(1)
    print("\nALL VERIFY OK: no dropped paragraphs in any book.")


if __name__ == '__main__':
    main()
