#!/usr/bin/env python3
"""Worker D: OT 담당 8권 Google Docs 업로드 (신규 생성) + API 검증 + export-back 검증.

Usage: python3 upload_ot_d.py Esther

게이트 7: 폴더 1I7rLOsUgbr5Ewyq7-IEshCk20scIBtPn에
'{한글이름} ({영어이름}) — MSG + Teen EN + KO' 제목으로 Doc 생성
-> docx 업로드(자동 변환) -> Docs API title 일치 + 테이블 수 == 장 수 확인
-> export-back 검증 (빈 틴 셀 없음, 누락 문단 없음).
결과는 gdocs_build/upload_results.json에 upsert + 본 스크립트 출력.
"""
import json, subprocess, sys, time, os

GDIR = '/home/hatch/workspace/teenz-bible-review/gdocs_build'
REVIEW = '/home/hatch/workspace/teenz-bible-review'
sys.path.insert(0, GDIR)
from build_gdocs import load_teen, parse_msg_txt, verify_no_drop

FOLDER = '1I7rLOsUgbr5Ewyq7-IEshCk20scIBtPn'

# (key, 한글, 영어, 장 수) — 작업자 D 담당 8권 (정경 순서)
BOOKS_D = [
    ('1Kings', '열왕기상', '1 Kings', 22),
    ('2Kings', '열왕기하', '2 Kings', 25),
    ('1Chronicles', '역대상', '1 Chronicles', 29),
    ('2Chronicles', '역대하', '2 Chronicles', 36),
    ('Ezra', '에스라', 'Ezra', 10),
    ('Nehemiah', '느헤미야', 'Nehemiah', 13),
    ('Esther', '에스더', 'Esther', 10),
    ('Isaiah', '이사야', 'Isaiah', 66),
]
BOOK_MAP = {k: (ko, en, n) for k, ko, en, n in BOOKS_D}

def run(args):
    r = subprocess.run(args, capture_output=True, text=True, timeout=300)
    if r.returncode != 0:
        raise RuntimeError(f'CMD FAILED: {" ".join(args[:6])}...\nSTDERR: {r.stderr[:600]}')
    out = r.stdout.strip()
    return json.loads(out) if out else {}

def verify_exported_ot(key, path):
    from docx import Document
    doc = Document(path)
    tables = doc.tables
    problems = []
    exp = BOOK_MAP[key][2]
    if len(tables) != exp:
        problems.append(f'table count {len(tables)} != {exp}')
    def cell_text(cell):
        return '\n'.join(p.text for p in cell.paragraphs).strip()
    for ti, t in enumerate(tables):
        rows = t.rows
        if not rows:
            problems.append(f'table {ti}: no rows'); continue
        for ri, row in enumerate(rows[1:], start=1):
            c0, c1, c2 = (cell_text(c) for c in row.cells)
            if not c0:
                if not c1 or not c2:
                    problems.append(f'table {ti} row {ri}: empty MSG cell but Teen cells also empty')
            else:
                if not c1: problems.append(f'table {ti} row {ri}: MSG present but Teen EN empty')
                if not c2: problems.append(f'table {ti} row {ri}: MSG present but Teen KO empty')
    en_chs = load_teen('en', key); ko_chs = load_teen('ko', key)
    units = parse_msg_txt(os.path.join(REVIEW, f'msg_{key}.txt'))
    for m in verify_no_drop(key, path, en_chs, ko_chs, units):
        problems.append(f'DROPPED {m[0]} ch{m[1]}: {m[2][:50]!r}')
    return problems

def main():
    keys = sys.argv[1:]
    assert keys and all(k in BOOK_MAP for k in keys), 'usage: upload_ot_d.py <OT-D Key...>'
    res_path = os.path.join(GDIR, 'upload_results.json')
    results = json.load(open(res_path, encoding='utf-8'))
    id_by_key = {r['key']: r for r in results}
    for key in keys:
        ko, en, exp_tables = BOOK_MAP[key]
        title = f'{ko} ({en}) — MSG + Teen EN + KO'
        docx = os.path.join(GDIR, f'{key}.docx')
        assert os.path.exists(docx) and os.path.getsize(docx) > 0, f'missing {docx}'
        print(f'[{key}] creating doc...', flush=True)
        c = run(['hatch_gws_cli', 'drive', 'files', 'create',
                 '--params', '{"ignoreDefaultVisibility":true}',
                 '--json', json.dumps({'name': title,
                                       'mimeType': 'application/vnd.google-apps.document',
                                       'parents': [FOLDER]})])
        doc_id = c['id']
        time.sleep(1)
        run(['hatch_gws_cli', 'drive', 'files', 'update',
             '--params', json.dumps({'fileId': doc_id}), '--upload', docx])
        time.sleep(3)
        g = run(['hatch_gws_cli', 'docs', 'documents', 'get',
                 '--params', json.dumps({'documentId': doc_id,
                                         'fields': 'title,body(content(table))'})])
        n_tables = sum(1 for x in g.get('body', {}).get('content', []) if 'table' in x)
        title_ok = (g.get('title') == title)
        print(f'[{key}] API: title_ok={title_ok} tables={n_tables}/{exp_tables}', flush=True)
        exp_path = f'/tmp/gverify_{key}.docx'
        r = subprocess.run(['hatch_gws_cli', 'drive', 'files', 'export',
                            '--params', json.dumps({'fileId': doc_id,
                                'mimeType': 'application/vnd.openxmlformats-officedocument.wordprocessingml.document'}),
                            '--output', exp_path],
                           capture_output=True, text=True, timeout=300)
        problems = []
        if r.returncode != 0:
            problems.append('export failed: ' + r.stderr[:200])
        else:
            problems = verify_exported_ot(key, exp_path)
        f = run(['hatch_gws_cli', 'drive', 'files', 'get',
                 '--params', json.dumps({'fileId': doc_id,
                                         'fields': 'id,name,webViewLink,parents'})])
        ok = title_ok and n_tables == exp_tables and not problems
        entry = {'key': key, 'title': title, 'id': doc_id,
                 'webViewLink': f.get('webViewLink'),
                 'tables': n_tables, 'expected': exp_tables,
                 'status': 'OK' if ok else 'MISMATCH',
                 'export_problems': problems[:10]}
        if key in id_by_key:
            results[results.index(id_by_key[key])] = entry
        else:
            results.append(entry)
        json.dump(results, open(res_path, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
        print(f'[{key}] {"VERIFIED OK" if ok else "PROBLEMS: " + str(problems[:6])}', flush=True)
        print(f'[{key}] {f.get("webViewLink")}', flush=True)
        if not ok:
            sys.exit(1)

if __name__ == '__main__':
    main()
