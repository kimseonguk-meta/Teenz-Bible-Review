"""Build the Matthew 3-way comparison docx: MSG vs Teenz EN vs Teenz KO."""
import json, re, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

from msg_parse import parse_msg_sections, range_to_set

WORK = '/home/hatch/workspace/teenz-bible-review'
OUT_DIR = os.path.expanduser('~/workspace/your_files/teenz-bible-review-matthew')
os.makedirs(OUT_DIR, exist_ok=True)

CHAPTER_NOTES = {
    1: 'EN은 족보(2–16절) 전체를 ‘3시대 요약’ 3개 문단으로 대체했습니다 (본문에 “이름 나열 대신 요약”이라고 명시). KO는 족보 이름을 전부 번역했습니다.',
}

KO_LABEL_RE = re.compile(r'^\d+\s*(~\s*\d+)?\s*절\.?$')

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:fill'), hexcolor)
    tcPr.append(shd)

def set_cell_text(cell, text, bold=False, size=9, color=None, italic=False):
    cell.text = ''
    p = cell.paragraphs[0]
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p

def add_para(cell, text, bold=False, size=9, color=None, italic=False, first=False):
    p = cell.paragraphs[0] if first else cell.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)
    return p

def load_msg():
    msg = {}
    raw_path = os.path.join(WORK, 'msg_matthew_raw.json')
    if os.path.exists(raw_path):
        msg.update(json.load(open(raw_path, encoding='utf-8')))
    for ch in ('1', '3'):
        p = os.path.join(WORK, f'msg_raw_{ch}.txt')
        if ch not in msg and os.path.exists(p):
            msg[ch] = open(p, encoding='utf-8').read()
    return msg

def main(chapters=None):
    en = json.load(open(os.path.join(WORK, 'en_matthew.json'), encoding='utf-8'))
    ko = json.load(open(os.path.join(WORK, 'ko_matthew.json'), encoding='utf-8'))
    msg = load_msg()
    # KJV verse counts for Matthew 1-28 (structural reference for MSG parsing guard)
    mt_verse_counts = [25,23,17,25,48,34,29,34,38,42,30,50,58,36,39,28,27,35,30,34,46,46,39,51,46,75,66,20]

    en_by = {c['num']: c for c in en}
    ko_by = {c['num']: c for c in ko}

    doc = Document()
    for section in doc.sections:
        section.orientation = 1  # landscape
        section.page_width, section.page_height = Cm(29.7), Cm(21.0)
        section.left_margin = section.right_margin = Cm(1.2)
        section.top_margin = section.bottom_margin = Cm(1.2)

    style = doc.styles['Normal']
    style.font.size = Pt(10)

    doc.add_heading('Teenz Bible 번역 대조 — 마태복음', level=1)
    doc.add_paragraph(
        'The Message 영어 원문(유진 피터슨) vs Teenz Bible 영어 번역 vs Teenz Bible 한국어 번역을 절 구간별로 나란히 정리한 검토용 문서입니다. '
        '각 행은 The Message의 절 구간을 기준으로 하며, 해당 구간에 대응하는 Teenz 문단이 없으면 “없음”으로 표시됩니다.'
    ).runs[0].font.size = Pt(10)
    doc.add_paragraph('작성: 2026-09-22 · 앱에는 미반영 (검토용)').runs[0].font.size = Pt(9)

    gap_summary = []

    ch_nums = chapters or sorted(set(en_by) & set(int(k) for k in msg.keys()))
    for n in ch_nums:
        e = en_by[n]
        k = ko_by[n]
        max_v = mt_verse_counts[n - 1]
        sections = parse_msg_sections(msg[str(n)], n, max_v)

        doc.add_heading(f'Matthew {n}' + (f' — {e["title"]}' if e.get('title') else ''), level=2)
        if n in CHAPTER_NOTES:
            p = doc.add_paragraph()
            r = p.add_run('※ ' + CHAPTER_NOTES[n])
            r.italic = True
            r.font.size = Pt(9)

        en_sets = [range_to_set(b) for b in (e['verseRanges'] or [])]
        ko_paras = []
        for ptxt, vr in zip(k['paragraphs'], k['verseRanges'] or []):
            if KO_LABEL_RE.match(ptxt.strip()):
                continue
            ko_paras.append((ptxt, range_to_set(vr)))
        en_paras = [(ptxt, s) for ptxt, s in zip(e['paragraphs'], en_sets)]

        table = doc.add_table(rows=1, cols=4)
        table.style = 'Table Grid'
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        hdr = table.rows[0].cells
        for i, t in enumerate(['절', 'The Message (원문)', 'Teenz Bible EN', 'Teenz Bible KO']):
            set_cell_text(hdr[i], t, bold=True, size=10, color=(255, 255, 255))
            shade(hdr[i], '1F3864')
        widths = [Cm(1.6), Cm(8.2), Cm(8.2), Cm(8.2)]
        for i, w in enumerate(widths):
            hdr[i].width = w

        used_en = set()
        for rng, header, mtxt in sections:
            mset = range_to_set(rng)
            en_hit = [(t, s) for idx, (t, s) in enumerate(en_paras) if s & mset]
            ko_hit = [(t, s) for (t, s) in ko_paras if s & mset]
            for idx, _ in enumerate(en_hit):
                used_en.add(en_paras.index(en_hit[idx]))
            row = table.add_row().cells
            set_cell_text(row[0], rng, bold=True, size=9)
            first = True
            if header:
                add_para(row[1], header, bold=True, size=9, first=True)
                first = False
            add_para(row[1], mtxt, size=9, first=first)
            if en_hit:
                f = True
                for t, s in en_hit:
                    is_hdr = t.startswith('§')
                    add_para(row[2], t[1:] if is_hdr else t, bold=is_hdr, size=9, first=f)
                    f = False
            else:
                set_cell_text(row[2], '⚠ 없음', size=9, color=(156, 0, 0), italic=True)
                shade(row[2], 'FFC7CE')
                gap_summary.append((n, rng, 'EN'))
            if ko_hit:
                f = True
                for t, s in ko_hit:
                    is_hdr = t.startswith('§')
                    add_para(row[3], t[1:] if is_hdr else t, bold=is_hdr, size=9, first=f)
                    f = False
            else:
                set_cell_text(row[3], '⚠ 없음', size=9, color=(156, 0, 0), italic=True)
                shade(row[3], 'FFC7CE')
                gap_summary.append((n, rng, 'KO'))

        # EN paragraphs that matched no MSG section
        for idx, (t, s) in enumerate(en_paras):
            if idx not in used_en and s:
                row = table.add_row().cells
                badge = (e['verseRanges'] or [])[idx]
                set_cell_text(row[0], str(badge) + ' (EN만)', bold=True, size=9)
                add_para(row[2], t[1:] if t.startswith('§') else t, bold=t.startswith('§'), size=9, first=True)
                set_cell_text(row[1], '—', size=9, color=(128, 128, 128))
                set_cell_text(row[3], '—', size=9, color=(128, 128, 128))

        doc.add_paragraph('')

    # gap summary at the end
    doc.add_heading('부록: 대응 문단이 없는 구간 목록', level=1)
    doc.add_paragraph('위 표에서 “⚠ 없음”으로 표시된 구간을 모은 목록입니다. (배지 오기입일 수도 있으니 원문과 함께 판단해 주세요.)').runs[0].font.size = Pt(9)
    if gap_summary:
        t2 = doc.add_table(rows=1, cols=3)
        t2.style = 'Table Grid'
        for i, h in enumerate(['장', '절 구간', '없는 쪽']):
            set_cell_text(t2.rows[0].cells[i], h, bold=True, size=10, color=(255, 255, 255))
            shade(t2.rows[0].cells[i], '1F3864')
        for n, rng, side in gap_summary:
            row = t2.add_row().cells
            set_cell_text(row[0], f'Matthew {n}', size=9)
            set_cell_text(row[1], rng, size=9)
            set_cell_text(row[2], side, size=9)
    else:
        doc.add_paragraph('없음 — 전 구간 대응됨.')

    out = os.path.join(OUT_DIR, 'Teenz-Bible-마태복음-번역대조.docx')
    doc.save(out)
    print('saved:', out, '| chapters:', ch_nums, '| gaps:', len(gap_summary))

if __name__ == '__main__':
    import sys
    chs = [int(x) for x in sys.argv[1:]] if len(sys.argv) > 1 else None
    main(chs)
